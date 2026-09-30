"""Excel export of a Risk Register: content, formatting and filename safety."""

from __future__ import annotations

import datetime as dt
from io import BytesIO

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from riskapp.excel_export import (
    HEADERS,
    risk_register_filename,
    safe_filename_part,
)

XLSX_MEDIA_TYPE = (
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
)


def _project(client: TestClient, name: str = "Project Alpha") -> dict:
    return client.post(
        "/api/projects",
        json={"name": name, "department": "D", "project_type": "T", "customer": "C"},
    ).json()


def _risk(client: TestClient, project_id: int, desc: str) -> dict:
    return client.post(
        "/api/risks",
        json={
            "project_id": project_id,
            "description": desc,
            "likelihood": "High",
            "impact": "Medium",
            "category": "Technical",
            "risk_start_date": "2026-09-01",
            "risk_end_date": "2026-10-01",
        },
    ).json()


class TestFilenameSafety:
    def test_sanitizes_special_characters(self) -> None:
        assert safe_filename_part("A/B: C*?<>|") == "A_B_C"
        assert safe_filename_part("") == "Register"

    def test_filename_shape(self) -> None:
        name = risk_register_filename("Project Alpha", dt.date(2026, 9, 30))
        assert name == "Risk_Register_Project_Alpha_2026-09-30.xlsx"


class TestExportEndpoint:
    def test_export_returns_a_formatted_workbook(self, client: TestClient) -> None:
        project = _project(client)
        _risk(client, project["id"], "First risk")
        _risk(client, project["id"], "Second risk")

        resp = client.get(f"/api/projects/{project['id']}/risks/export")
        assert resp.status_code == 200
        assert resp.headers["content-type"].startswith(XLSX_MEDIA_TYPE)
        assert (
            'filename="Risk_Register_Project_Alpha_'
            in resp.headers["content-disposition"]
        )

        workbook = load_workbook(BytesIO(resp.content))
        sheet = workbook.active
        header = [cell.value for cell in sheet[1]]
        assert header == list(HEADERS)
        assert sheet.freeze_panes == "A2"
        assert sheet.auto_filter.ref is not None
        # Header row + two data rows.
        assert sheet.max_row == 3
        assert {sheet.cell(row=r, column=2).value for r in (2, 3)} == {
            "First risk",
            "Second risk",
        }

    def test_export_includes_owner_details_and_acknowledgement(
        self, client: TestClient
    ) -> None:
        project = _project(client)
        risk = _risk(client, project["id"], "Owned risk")
        owner = client.post(
            "/api/external-owners",
            json={
                "full_name": "Jane Smith",
                "email": "jane.smith@abc.com",
                "organization": "ABC",
            },
        ).json()
        client.patch(f"/api/risks/{risk['id']}", json={"owner_user_id": owner["id"]})

        resp = client.get(f"/api/projects/{project['id']}/risks/export")
        sheet = load_workbook(BytesIO(resp.content)).active
        row = [cell.value for cell in sheet[2]]
        assert row[HEADERS.index("Risk Owner")] == "Jane Smith"
        assert row[HEADERS.index("Risk Owner Email")] == "jane.smith@abc.com"
        assert row[HEADERS.index("Owner Type")] == "External"
        assert row[HEADERS.index("Acknowledged")] == "Not Acknowledged"

    def test_export_only_contains_this_register(self, client: TestClient) -> None:
        project_a = _project(client, "Alpha")
        project_b = _project(client, "Beta")
        _risk(client, project_a["id"], "Only Alpha")
        _risk(client, project_b["id"], "Only Beta")

        resp = client.get(f"/api/projects/{project_a['id']}/risks/export")
        sheet = load_workbook(BytesIO(resp.content)).active
        descriptions = {sheet.cell(row=r, column=2).value for r in range(2, sheet.max_row + 1)}
        assert descriptions == {"Only Alpha"}
