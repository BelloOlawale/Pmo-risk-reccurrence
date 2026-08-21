"""Tests for the Excel parser — read .xlsx, detect headers, extract rows."""

from __future__ import annotations

from pathlib import Path

import openpyxl

from riskapp.import_pipeline.excel_parser import parse_excel

FIXTURES_DIR = Path(__file__).parent / "fixtures"


class TestParseExcel:
    """Integration tests against real fixture files."""

    def test_parses_punuka_file_returns_12_rows(self) -> None:
        rows = parse_excel(str(FIXTURES_DIR / "punuka_bpa.xlsx"))
        assert len(rows) == 12

    def test_parsed_rows_are_dicts_with_header_keys(self) -> None:
        rows = parse_excel(str(FIXTURES_DIR / "punuka_bpa.xlsx"))
        first = rows[0]
        assert isinstance(first, dict)
        assert "Risk ID" in first
        assert "Risk Description" in first
        assert "Risk Category" in first
        assert first["Risk ID"] == "R001"

    def test_parses_wacl_file_returns_7_rows(self) -> None:
        rows = parse_excel(str(FIXTURES_DIR / "wacl_bpa.xlsx"))
        assert len(rows) == 7
        first = rows[0]
        assert "ID" in first
        assert "Description of Risk" in first
        assert "Risk Score" in first
        assert first["ID"] == "1"

    def test_parses_tnl_file_with_hml_format(self) -> None:
        rows = parse_excel(str(FIXTURES_DIR / "tnl_erp.xlsx"))
        assert len(rows) == 5
        assert "Probability (H/M/L)" in rows[0]
        assert rows[0]["Probability (H/M/L)"] == "H"

    def test_parses_seamless_hr_with_blank_column_a(self) -> None:
        rows = parse_excel(str(FIXTURES_DIR / "seamless_hr_ai.xlsx"))
        assert len(rows) == 4
        assert "S/N" in rows[0]
        assert "Risk Description" in rows[0]
        assert rows[0]["S/N"] == "1"

    def test_header_detection_scores_rows_not_first_keyword_match(
        self, tmp_path: Path
    ) -> None:
        """Scoring ensures the real multi-column header row wins over metadata
        rows containing incidental keyword hits."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.cell(row=1, column=2, value="Risk Scoring Criteria :")
        ws.cell(row=2, column=2, value="Impact (I): High, Medium, Low")
        ws.cell(row=3, column=2, value="Probability (P): High, Medium, Low")
        ws.cell(row=5, column=1, value="ID")
        ws.cell(row=5, column=2, value="Risk Description")
        ws.cell(row=5, column=3, value="Category")
        ws.cell(row=5, column=4, value="Impact")
        ws.cell(row=5, column=5, value="Probability")
        ws.cell(row=5, column=6, value="Risk Score")
        ws.cell(row=5, column=7, value="Owner")
        ws.cell(row=6, column=1, value="R001")
        ws.cell(row=6, column=2, value="Scope creep")
        ws.cell(row=6, column=3, value="Scope")
        ws.cell(row=6, column=4, value="High")
        ws.cell(row=6, column=5, value="Medium")

        path = tmp_path / "header.xlsx"
        wb.save(str(path))
        wb.close()

        rows = parse_excel(str(path))
        assert len(rows) == 1
        assert rows[0]["ID"] == "R001"
        assert rows[0]["Risk Description"] == "Scope creep"
        assert rows[0]["Owner"] == ""

    def test_parsing_ole2_xls_file_returns_empty(self, tmp_path: Path) -> None:
        """Files that aren't valid .xlsx (OLE2 .xls signature) return empty."""
        path = tmp_path / "fake.xlsx"
        path.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 100)
        assert parse_excel(str(path)) == []

    def test_missing_file_returns_empty(self) -> None:
        assert parse_excel(str(FIXTURES_DIR / "does-not-exist.xlsx")) == []
