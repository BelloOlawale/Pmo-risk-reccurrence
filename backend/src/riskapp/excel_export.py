"""Excel (.xlsx) export of a project's Risk Register.

Uses openpyxl (already a project dependency) to render the current database
rows — never mocked data — into a professionally formatted workbook: a frozen,
filterable header row, sensible column widths, and consistent date formatting.
"""

from __future__ import annotations

import datetime as dt
import re
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from riskapp import models
from riskapp.domain.sla import as_naive_utc
from riskapp.external import EXTERNAL  # noqa: F401  (documents the Owner Type values)

_DATE_FMT = "yyyy-mm-dd"
_DATETIME_FMT = "yyyy-mm-dd hh:mm"

# Column order and headers for the export. Kept as one list so the header row
# and the row builder can never drift apart.
HEADERS: tuple[str, ...] = (
    "Risk ID",
    "Risk Title / Description",
    "Project",
    "Risk Owner",
    "Risk Owner Email",
    "Risk Owner Type",
    "Risk Category",
    "Likelihood",
    "Impact",
    "Severity / Risk Rating",
    "Response Strategy",
    "Project Lifecycle",
    "Risk Start Date",
    "Risk End Date",
    "SLA Start Date",
    "SLA Deadline",
    "Status",
    "Acknowledged",
    "Acknowledgement Date",
    "Created Date",
    "Updated Date",
)

# 1-based column indexes that hold dates vs datetimes, for number formatting.
_DATE_COLUMNS = (13, 14, 15)
_DATETIME_COLUMNS = (16, 19, 20, 21)


def safe_filename_part(name: str | None) -> str:
    """Reduce a project name to something safe for a Windows/POSIX filename."""
    cleaned = re.sub(r"[^A-Za-z0-9._ -]+", "_", name or "").strip()
    cleaned = re.sub(r"[\s-]+", "_", cleaned)
    cleaned = re.sub(r"_+", "_", cleaned).strip("._")
    return cleaned or "Register"


def risk_register_filename(project_name: str | None, on: dt.date | None = None) -> str:
    """``Risk_Register_<Project>_<YYYY-MM-DD>.xlsx`` with a sanitized project name."""
    date = on or dt.date.today()
    return f"Risk_Register_{safe_filename_part(project_name)}_{date.isoformat()}.xlsx"


def _date_value(value: dt.date | None) -> dt.date | str:
    return value if value else ""


def _datetime_value(value: dt.datetime | None) -> dt.datetime | str:
    # openpyxl rejects timezone-aware datetimes; the app stores naive UTC, but
    # a Postgres TIMESTAMPTZ column can hand one back, so normalise defensively.
    return as_naive_utc(value) if value else ""


def _row(project: models.Project, risk: models.Risk) -> list[Any]:
    owner = risk.owner
    return [
        risk.risk_code,
        risk.description,
        project.name,
        owner.display_name if owner else "",
        owner.upn if owner else "",
        (owner.owner_type if owner else ""),
        risk.category or "",
        risk.likelihood,
        risk.impact,
        risk.risk_rating,
        risk.response_strategy or "",
        risk.identified_during or "",
        _date_value(risk.risk_start_date),
        _date_value(risk.risk_end_date),
        # The Risk Start Date is the SLA start.
        _date_value(risk.risk_start_date),
        _datetime_value(risk.sla_deadline),
        risk.status,
        "Acknowledged" if risk.sla_acknowledged else "Not Acknowledged",
        _datetime_value(risk.acknowledged_at),
        _datetime_value(risk.created_at),
        _datetime_value(risk.updated_at),
    ]


def _auto_widths(ws: Any, rows: list[list[Any]]) -> None:
    """Size each column to its content (bounded), so the file opens readable."""
    for col_idx, header in enumerate(HEADERS, start=1):
        longest = len(str(header))
        for row in rows:
            value = row[col_idx - 1]
            if value is None:
                continue
            longest = max(longest, len(str(value)))
        width = min(max(longest + 2, 12), 60)
        ws.column_dimensions[get_column_letter(col_idx)].width = width


def build_risk_register_workbook(
    project: models.Project, risks: list[models.Risk]
) -> bytes:
    """Render the register to xlsx bytes (header frozen + filterable)."""
    rows = [_row(project, risk) for risk in risks]

    workbook = Workbook()
    workbook.properties.title = "WRAGBY RiskIntel — Risk Register"
    workbook.properties.subject = "PMO Risk Management"
    workbook.properties.creator = "WRAGBY RiskIntel"
    sheet = workbook.active
    sheet.title = "Risk Register"
    sheet.append(list(HEADERS))
    for row in rows:
        sheet.append(row)

    header_fill = PatternFill("solid", fgColor="ED1C2E")
    header_font = Font(bold=True, color="FFFFFF")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(vertical="center", wrap_text=True)

    _auto_widths(sheet, rows)

    for col_idx in _DATE_COLUMNS:
        for row_idx in range(2, sheet.max_row + 1):
            sheet.cell(row=row_idx, column=col_idx).number_format = _DATE_FMT
    for col_idx in _DATETIME_COLUMNS:
        for row_idx in range(2, sheet.max_row + 1):
            cell = sheet.cell(row=row_idx, column=col_idx)
            if cell.value:
                cell.number_format = _DATETIME_FMT

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
